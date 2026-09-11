from pathlib import Path
import re,json,html
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parent;OUT=ROOT.parent/'outputs';OUT.mkdir(exist_ok=True)
raw=(ROOT/'manuscript_source.md').read_text(encoding='utf-8')
refs={r['id']:r for r in json.loads((ROOT/'literature/references.json').read_text(encoding='utf-8'))};order=[]
def cite(m):
 key=m.group(1);assert key in refs and refs[key]['metadata_status']=='crossref_verified'
 if key not in order:order.append(key)
 return f'[{order.index(key)+1}](https://doi.org/{refs[key]["doi"]})'
md=re.sub(r'\[@(\w+)\]',cite,raw)
bib=[]
for i,key in enumerate(order,1):
 r=refs[key];authors=', '.join(r['authors'][:6])+(', et al.' if len(r['authors'])>6 else '.')
 article=r.get('pages') or r.get('article_number') or ''
 def clean(s):return ' '.join(html.unescape(re.sub(r'<[^>]+>','',s)).split()).replace('α v β 3','αvβ3')
 bib.append(f'{i}. {clean(authors)} {clean(r["title"])}. {clean(r["journal"])}. {r["year"]};{r.get("volume") or ""}:{article}. [doi:{r["doi"]}](https://doi.org/{r["doi"]})')
md=md.replace('<!-- REFERENCES_GENERATED -->','\n\n'.join(bib))
(OUT/'manuscript_zh.md').write_text(md,encoding='utf-8')
(OUT/'references.json').write_text(json.dumps([refs[k] for k in order],ensure_ascii=False,indent=2),encoding='utf-8')
doc=Document()
for e in doc.styles.element.xpath('.//w:pBdr'):e.getparent().remove(e)
section=doc.sections[0];section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=Inches(.72);section.bottom_margin=Inches(.72);section.left_margin=Inches(.65);section.right_margin=Inches(.65)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0);st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'SimSun')
normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.3;normal.paragraph_format.space_after=Pt(6)
for name,size in [('Title',17),('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
 st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(6);st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'SimHei')
doc.styles['Title'].paragraph_format.space_before=Pt(0)
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');cached=OxmlElement('w:r');cachedtext=OxmlElement('w:t');cachedtext.text='1';cached.append(cachedtext);fld.append(cached);footer._p.append(fld)
def inline(p,text):
 for part in re.split(r'(\*\*.*?\*\*|\[[^\]]+\]\([^)]+\))',text):
  if not part:continue
  if part.startswith('**') and part.endswith('**'):p.add_run(part[2:-2]).bold=True
  elif re.fullmatch(r'\[[^\]]+\]\([^)]+\)',part):
   m=re.match(r'\[([^\]]+)\]\(([^)]+)\)',part);label,url=m.groups();h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'000000');pr.append(color);r.append(pr);t=OxmlElement('w:t');t.text='['+label+']' if label.isdigit() else label;r.append(t);h.append(r);p._p.append(h)
  else:p.add_run(part)
def table(lines):
 rows=[[c.strip() for c in l.strip().strip('|').split('|')] for l in lines if not re.match(r'^\|[\s:|\-]+\|$',l)]
 t=doc.add_table(rows=1,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=[1.2,2.55,1.65,1.8] if len(rows[0])==4 else [2.65,1,1,1.35,1.2]
 for i,col in enumerate(t.columns):col.width=Inches(widths[i])
 for i,c in enumerate(t.rows[0].cells):c.width=Inches(widths[i])
 for ri,row in enumerate(rows):
  cells=t.rows[0].cells if ri==0 else t.add_row().cells
  for ci,txt in enumerate(row):
   cells[ci].width=Inches(widths[ci]);cells[ci].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;p=cells[ci].paragraphs[0];p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);inline(p,txt)
   for r in p.runs:r.font.size=Pt(9);r.bold=(ri==0)
   tcpr=cells[ci]._tc.get_or_add_tcPr();border=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');border.append(e)
   tcpr.append(border)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E6EEF3');tcpr.append(sh)
  trpr=t.rows[ri]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if ri==0:rep=OxmlElement('w:tblHeader');trpr.append(rep)
 doc.add_paragraph()
lines=md.splitlines();i=0
while i<len(lines):
 l=lines[i].strip();i+=1
 if not l:continue
 if l.startswith('|'):
  block=[l]
  while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i]);i+=1
  table(block);continue
 if l.startswith('# '):p=doc.add_paragraph(style='Title');inline(p,l[2:]);continue
 if l.startswith('## '):
  p=doc.add_paragraph(style='Heading 1');inline(p,l[3:]);continue
 if l.startswith('### '):p=doc.add_paragraph(style='Heading 2');inline(p,l[4:]);continue
 p=doc.add_paragraph();inline(p,l)
 if l.startswith('Cross-cohort'):p.paragraph_format.space_after=Pt(8)
# Figures are consolidated after legends/references, with explicit full-size figure pages.
for f in sorted((OUT/'figures').glob('*.png')):
 doc.add_page_break();p=doc.add_paragraph(style='Heading 1');p.add_run(f.stem.split('_')[0].replace('FigureS','Supplementary Figure S').replace('Figure','Figure '))
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(f),width=Inches(7.2))
doc.core_properties.title=raw.splitlines()[0][2:];doc.core_properties.subject='公开数据重分析；探索性跨队列研究';doc.core_properties.author='';doc.core_properties.keywords='cerebral ischemia; astrocyte; endothelial; public data'
dest=OUT/'manuscript_zh.docx';doc.save(dest)
print(json.dumps({'docx':str(dest),'references':len(order),'markdown_characters':len(md),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables)},ensure_ascii=False))
